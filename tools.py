"""All of Master's tools live here.

To add a tool:
  1. write the function
  2. add its description to TOOLS
  3. register it in TOOL_FUNCTIONS
"""

import io
import re
import json
import os
import glob
import math
import ctypes
import tempfile
import threading
import subprocess
import webbrowser
from datetime import datetime, timedelta

import requests
from ddgs import DDGS
import win32com.client
from docx import Document as _Docx
from pptx import Presentation as _Pptx
from pptx.util import Inches as _In, Pt as _Pt
from pptx.dml.color import RGBColor as _RGB
from pptx.enum.text import PP_ALIGN as _ALIGN
from pptx.enum.shapes import MSO_SHAPE as _SHAPE
from openpyxl import Workbook as _Workbook, load_workbook as _load_wb
from pypdf import PdfReader as _PdfReader

TASKS_FILE = "tasks.json"
MEMORY_FILE = "memory.json"
REMINDERS_FILE = "reminders.json"
CONTACTS_FILE = "contacts.json"

# Master_code.py sets these to its speak() function
on_timer = None      # for set_timer
announce = None       # for reminders firing in the background


# ---------- info tools ----------

def get_time():
    """Return the current date and time as readable text."""
    return datetime.now().strftime("%A %d %B %Y, %I:%M %p")


def web_search(query, num_results=5):
    """Search the web and return the top results as text."""
    try:
        results = DDGS().text(query, max_results=num_results)
    except Exception as e:
        return f"Search failed: {e}"
    if not results:
        return "No results found."
    lines = []
    for r in results:
        lines.append(f"- {r['title']}\n  {r['body']}\n  ({r['href']})")
    return "\n".join(lines)


_WMO = {
    0: "clear sky", 1: "mainly clear", 2: "partly cloudy", 3: "overcast",
    45: "fog", 48: "freezing fog",
    51: "light drizzle", 53: "drizzle", 55: "heavy drizzle",
    56: "freezing drizzle", 57: "heavy freezing drizzle",
    61: "light rain", 63: "rain", 65: "heavy rain",
    66: "freezing rain", 67: "heavy freezing rain",
    71: "light snow", 73: "snow", 75: "heavy snow", 77: "snow grains",
    80: "light rain showers", 81: "rain showers", 82: "violent rain showers",
    85: "light snow showers", 86: "snow showers",
    95: "thunderstorm", 96: "thunderstorm with hail", 99: "severe thunderstorm with hail",
}


def get_weather(place):
    """Current weather for a city or place name. No API key needed (open-meteo)."""
    try:
        g = requests.get("https://geocoding-api.open-meteo.com/v1/search",
                         params={"name": place, "count": 1}, timeout=15).json()
        if not g.get("results"):
            return f"Couldn't find a place called {place}."
        loc = g["results"][0]
        name = loc["name"] + (f", {loc['country']}" if loc.get("country") else "")
        w = requests.get(
            "https://api.open-meteo.com/v1/forecast",
            params={"latitude": loc["latitude"], "longitude": loc["longitude"],
                    "current": "temperature_2m,apparent_temperature,weather_code,wind_speed_10m"},
            timeout=15).json()
        c = w["current"]
        desc = _WMO.get(c["weather_code"], "unknown conditions")
        return (f"{name}: {desc}, {round(c['temperature_2m'])} degrees "
                f"(feels like {round(c['apparent_temperature'])}), "
                f"wind {round(c['wind_speed_10m'])} kilometres per hour.")
    except Exception as e:
        return f"Weather lookup failed: {e}"


# ---------- action tools ----------

def open_website(url):
    """Open a web page in the default browser."""
    if not re.match(r"^https?://", url):
        url = "https://" + url
    try:
        webbrowser.open(url)
        return f"Opening {url}"
    except Exception as e:
        return f"Couldn't open {url}: {e}"


def _start_apps():
    """Every app in the Start menu as (name, appid) - includes Store/UWP and Office apps."""
    try:
        out = subprocess.run(
            ["powershell", "-NoProfile", "-Command", "Get-StartApps | ConvertTo-Json -Compress"],
            capture_output=True, text=True, timeout=15).stdout
        data = json.loads(out)
        if isinstance(data, dict):
            data = [data]
        return [(d["Name"], d["AppID"]) for d in data]
    except Exception:
        return []


_MEDIA_VK = {
    "play_pause": 0xB3, "next": 0xB0, "previous": 0xB1, "stop": 0xB2,
    "volume_up": 0xAF, "volume_down": 0xAE, "mute": 0xAD,
}


def media_control(action):
    """Media / volume keys: play_pause, next, previous, stop, volume_up, volume_down, mute."""
    vk = _MEDIA_VK.get(action)
    if vk is None:
        return f"Unknown media action: {action}"
    reps = 5 if action in ("volume_up", "volume_down") else 1
    for _ in range(reps):
        ctypes.windll.user32.keybd_event(vk, 0, 0, 0)
        ctypes.windll.user32.keybd_event(vk, 0, 2, 0)     # key up
    return f"Done: {action.replace('_', ' ')}."


def system_control(action):
    """Lock or sleep the computer. action is 'lock' or 'sleep'."""
    if action == "lock":
        ctypes.windll.user32.LockWorkStation()
        return "Locking the screen."
    if action == "sleep":
        subprocess.Popen(["rundll32.exe", "powrprof.dll,SetSuspendState", "0,1,0"])
        return "Going to sleep."
    return f"Unknown system action: {action}"


def set_timer(minutes, label=""):
    """Set a timer for a number of minutes; it announces itself when it finishes."""
    try:
        secs = float(minutes) * 60
    except (TypeError, ValueError):
        return f"Couldn't understand '{minutes}' minutes."
    if secs <= 0:
        return "The timer needs to be more than zero minutes."

    def done():
        msg = "Timer finished" + (f": {label}." if label else ".")
        if on_timer:
            on_timer(msg)

    threading.Timer(secs, done).start()
    unit = "minute" if float(minutes) == 1 else "minutes"
    return f"Timer set for {minutes} {unit}" + (f" for {label}." if label else ".")


def open_app(name):
    """Open an application by name. Matches Start-menu apps including Microsoft Store and Office apps."""
    want = name.strip().lower()
    apps = _start_apps()
    match = None
    tests = (
        lambda n: n == want,
        lambda n: n.startswith(want),
        lambda n: want in n,
        lambda n: len(n) >= 4 and n in want,    # "word" matches spoken "microsoft word"
    )
    for test in tests:
        for n, aid in apps:
            if test(n.lower()):
                match = (n, aid)
                break
        if match:
            break
    try:
        if match:
            subprocess.Popen(["explorer.exe", f"shell:AppsFolder\\{match[1]}"])
            return f"Opening {match[0]}."
        subprocess.Popen(["cmd", "/c", "start", "", name])   # fallback for non-Start-menu programs
        return f"Opening {name}."
    except Exception as e:
        return f"Couldn't open {name}: {e}"


# ---------- office documents ----------

def _docs_dir():
    d = os.path.join(os.path.expanduser("~"), "Documents", "Master")
    os.makedirs(d, exist_ok=True)
    return d


def _new_path(title, ext):
    base = re.sub(r"[^\w -]", "", title or "").strip() or "document"
    stamp = datetime.now().strftime("%Y-%m-%d %H%M")
    return os.path.join(_docs_dir(), f"{base} {stamp}.{ext}")


def create_word_document(title, content):
    """Create a Word .docx from generated text and open it.
    In `content`: blank lines separate paragraphs; a line starting with '# ' is a heading,
    '## ' a sub-heading, '- ' a bullet."""
    try:
        doc = _Docx()
        doc.add_heading(title, level=0)
        for block in content.split("\n\n"):
            block = block.strip()
            if not block:
                continue
            if block.startswith("## "):
                doc.add_heading(block[3:].strip(), level=2)
            elif block.startswith("# "):
                doc.add_heading(block[2:].strip(), level=1)
            elif block.startswith("- "):
                for line in block.splitlines():
                    doc.add_paragraph(line.lstrip("- ").strip(), style="List Bullet")
            else:
                doc.add_paragraph(block)
        path = _new_path(title, "docx")
        doc.save(path)
        os.startfile(path)
        return f"Created and opened {os.path.basename(path)}."
    except Exception as e:
        return f"Couldn't create the document: {e}"


def _fetch_image(query):
    """Search the web for an image, download and normalise it, return a PNG path (or None)."""
    try:
        from PIL import Image
        results = DDGS().images(query, max_results=6)
    except Exception:
        return None
    for r in results or []:
        url = r.get("image") or r.get("thumbnail")
        if not url:
            continue
        try:
            resp = requests.get(url, timeout=8, headers={"User-Agent": "Mozilla/5.0"})
            if resp.status_code != 200 or len(resp.content) < 2000:
                continue
            im = Image.open(io.BytesIO(resp.content)).convert("RGB")
            if im.width < 200 or im.height < 150:
                continue
            im.thumbnail((1400, 1400))
            path = os.path.join(tempfile.gettempdir(), f"master_img_{abs(hash(url)) % 10**8}.png")
            im.save(path, "PNG")
            return path
        except Exception:
            continue
    return None


# deck styling
_PPT_PRIMARY = _RGB(0x1B, 0x3A, 0x6B)   # deep blue - titles
_PPT_ACCENT = _RGB(0x2F, 0x6D, 0xF6)    # bright blue - bars / bullets
_PPT_TEXT = _RGB(0x23, 0x28, 0x2E)      # near-black body text
_PPT_MUTED = _RGB(0x8A, 0x93, 0xA2)     # grey footers
_PPT_FONT = "Segoe UI"


def _pw_para(tf, text, size, *, bold=False, color=_PPT_TEXT, align=_ALIGN.LEFT,
             spacing=1.15, after=8, first=False):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.text = text
    p.alignment = align
    p.line_spacing = spacing
    p.space_after = _Pt(after)
    f = p.font
    f.name = _PPT_FONT
    f.size = _Pt(size)
    f.bold = bold
    f.color.rgb = color
    return p


def _pw_bar(slide, left, top, w, h, color):
    sh = slide.shapes.add_shape(_SHAPE.RECTANGLE, left, top, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = color
    sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def create_powerpoint(title, slides):
    """Create a styled widescreen .pptx and open it. In `slides`: each slide is a title
    line, then '- ' bullet lines; separate slides with a blank line. A line
    'image: <what to show>' adds a web image."""
    try:
        prs = _Pptx()
        prs.slide_width = _In(13.333)
        prs.slide_height = _In(7.5)
        blank = prs.slide_layouts[6]

        # ---- title slide ----
        s = prs.slides.add_slide(blank)
        _pw_bar(s, 0, _In(3.02), _In(13.333), _In(0.07), _PPT_ACCENT)
        t = s.shapes.add_textbox(_In(0.9), _In(2.15), _In(11.5), _In(1.5)).text_frame
        t.word_wrap = True
        _pw_para(t, title, 40, bold=True, color=_PPT_PRIMARY, first=True)
        d = s.shapes.add_textbox(_In(0.92), _In(3.35), _In(11.5), _In(0.5)).text_frame
        _pw_para(d, datetime.now().strftime("%d %B %Y"), 15, color=_PPT_MUTED, first=True)

        # ---- content slides ----
        blocks = [b for b in slides.split("\n\n") if b.strip()]
        for n, block in enumerate(blocks, 1):
            lines = [ln.rstrip() for ln in block.splitlines() if ln.strip()]
            if not lines:
                continue

            img_query, bullets = None, []
            for ln in lines[1:]:
                if ln.strip().lower().startswith("image:"):
                    img_query = ln.split(":", 1)[1].strip()
                else:
                    b = ln.lstrip("-*• ").strip()
                    if b:
                        bullets.append(b)

            sl = prs.slides.add_slide(blank)

            ht = sl.shapes.add_textbox(_In(0.75), _In(0.5), _In(11.8), _In(1.0)).text_frame
            ht.word_wrap = True
            _pw_para(ht, lines[0].lstrip("# ").strip(), 29, bold=True,
                     color=_PPT_PRIMARY, first=True)
            _pw_bar(sl, _In(0.8), _In(1.42), _In(2.1), _In(0.05), _PPT_ACCENT)

            img_path = _fetch_image(img_query) if img_query else None
            body_w = _In(7.0) if img_path else _In(11.8)

            bf = sl.shapes.add_textbox(_In(0.8), _In(1.85), body_w, _In(4.9)).text_frame
            bf.word_wrap = True
            for i, b in enumerate(bullets):
                _pw_para(bf, "•  " + b, 19, color=_PPT_TEXT,
                         spacing=1.2, after=12, first=(i == 0))

            if img_path:
                try:
                    pic = sl.shapes.add_picture(img_path, _In(8.2), _In(1.85), width=_In(4.4))
                    pic.line.color.rgb = _RGB(0xD8, 0xDD, 0xE4)
                    pic.line.width = _Pt(1)
                except Exception:
                    pass

            ff = sl.shapes.add_textbox(_In(0.75), _In(7.0), _In(10), _In(0.4)).text_frame
            _pw_para(ff, title, 9, color=_PPT_MUTED, first=True)
            pf = sl.shapes.add_textbox(_In(12.35), _In(7.0), _In(0.6), _In(0.4)).text_frame
            _pw_para(pf, str(n), 9, color=_PPT_MUTED, align=_ALIGN.RIGHT, first=True)

        path = _new_path(title, "pptx")
        prs.save(path)
        os.startfile(path)
        return f"Created and opened {os.path.basename(path)}."
    except Exception as e:
        return f"Couldn't create the presentation: {e}"


def create_spreadsheet(title, data):
    """Create an .xlsx and open it. In `data`: one row per line, cells separated by | or comma."""
    try:
        wb = _Workbook()
        ws = wb.active
        ws.title = (title or "Sheet1")[:31]
        for line in data.splitlines():
            line = line.strip()
            if not line:
                continue
            ws.append(re.split(r"\s*[|,]\s*", line))
        path = _new_path(title, "xlsx")
        wb.save(path)
        os.startfile(path)
        return f"Created and opened {os.path.basename(path)}."
    except Exception as e:
        return f"Couldn't create the spreadsheet: {e}"


# ---------- email + calendar ----------

def _parse_dt(s):
    for fmt in ("%Y-%m-%d %H:%M", "%Y-%m-%d %I:%M %p", "%d %B %Y %H:%M", "%m/%d/%Y %H:%M"):
        try:
            return datetime.strptime((s or "").strip(), fmt)
        except (ValueError, AttributeError):
            continue
    return None


def _with_timeout(fn, seconds, on_timeout):
    """Run fn() in a worker thread; return on_timeout if it doesn't finish in time."""
    box = [on_timeout]

    def run():
        try:
            box[0] = fn()
        except Exception as e:
            box[0] = f"error: {e}"

    t = threading.Thread(target=run, daemon=True)
    t.start()
    t.join(seconds)
    return box[0]


def draft_email(to, subject, body):
    """Open a pre-filled draft in the default mail app for the user to review and send. Never sends."""
    from urllib.parse import quote
    try:
        url = (f"mailto:{to or ''}"
               f"?subject={quote(subject or '')}"
               f"&body={quote(body or '')}")
        os.startfile(url)
        return "Opened an email draft for you to review and send."
    except Exception as e:
        return f"Couldn't open an email draft: {e}"


def add_calendar_event(title, start, minutes=60):
    """Create a calendar invite (.ics) and open it so the user can add it to their calendar."""
    dt = _parse_dt(start)
    if dt is None:
        return f"Couldn't understand the date/time '{start}'. Use YYYY-MM-DD HH:MM."
    try:
        end = dt + timedelta(minutes=int(minutes))
        uid = datetime.now().strftime("%Y%m%d%H%M%S") + "@master"
        ics = (
            "BEGIN:VCALENDAR\r\nVERSION:2.0\r\nPRODID:-//Master//EN\r\nCALSCALE:GREGORIAN\r\n"
            "BEGIN:VEVENT\r\n"
            f"UID:{uid}\r\n"
            f"DTSTAMP:{datetime.now().strftime('%Y%m%dT%H%M%S')}\r\n"
            f"DTSTART:{dt.strftime('%Y%m%dT%H%M%S')}\r\n"
            f"DTEND:{end.strftime('%Y%m%dT%H%M%S')}\r\n"
            f"SUMMARY:{title}\r\n"
            "END:VEVENT\r\nEND:VCALENDAR\r\n"
        )
        path = _new_path(title, "ics")
        with open(path, "w", encoding="utf-8", newline="") as f:
            f.write(ics)
        os.startfile(path)
        return (f"Opened a calendar invite for {dt.strftime('%a %d %b %I:%M %p')} - "
                "accept it to add it to your calendar.")
    except Exception as e:
        return f"Couldn't create the event: {e}"


def _read_calendar(days):
    import pythoncom
    pythoncom.CoInitialize()
    try:
        ns = win32com.client.Dispatch("Outlook.Application").GetNamespace("MAPI")
        items = ns.GetDefaultFolder(9).Items            # 9 = calendar
        items.IncludeRecurrences = True
        items.Sort("[Start]")
        now = datetime.now()
        end = now + timedelta(days=days)
        rng = items.Restrict(
            "[Start] >= '%s' AND [Start] <= '%s'"
            % (now.strftime("%m/%d/%Y %H:%M"), end.strftime("%m/%d/%Y %H:%M")))
        out = []
        for it in rng:
            try:
                out.append(f"{it.Start.strftime('%a %d %b %I:%M %p')} - {it.Subject}")
            except Exception:
                continue
            if len(out) >= 25:
                break
        return "\n".join(out) if out else f"Nothing on the calendar in the next {days} days."
    finally:
        pythoncom.CoUninitialize()


def list_calendar_events(days=7):
    """List upcoming Outlook calendar events for the next N days (needs classic Outlook set up)."""
    return _with_timeout(
        lambda: _read_calendar(int(days)), 8,
        "Couldn't reach the calendar - classic Outlook needs to be open for this.")


# ---------- long-term memory ----------

def _load_json(path, default):
    if os.path.exists(path):
        try:
            with open(path, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception:
            return default
    return default


def _save_json(path, data):
    with open(path, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=2, ensure_ascii=False)


def remember(fact):
    """Store a fact about the user for the long term."""
    mem = _load_json(MEMORY_FILE, [])
    mem.append({"text": fact, "added": datetime.now().strftime("%Y-%m-%d")})
    _save_json(MEMORY_FILE, mem)
    return f"Noted: {fact}"


def recall(topic=""):
    """Recall stored facts. With a topic, return the ones that best match; without, return all."""
    mem = _load_json(MEMORY_FILE, [])
    if not mem:
        return "I don't have anything remembered yet."
    words = [w for w in re.findall(r"[a-z0-9]+", topic.lower()) if len(w) > 2]
    if not words:
        return "\n".join(f"- {m['text']}" for m in mem)
    scored = []
    for m in mem:
        hits = sum(1 for w in words if w in m["text"].lower())
        if hits:
            scored.append((hits, m["text"]))
    if not scored:
        return f"Nothing remembered about '{topic}'."
    scored.sort(reverse=True)
    return "\n".join(f"- {t}" for _, t in scored[:10])


def forget(topic):
    """Remove remembered facts that match the topic."""
    mem = _load_json(MEMORY_FILE, [])
    words = [w for w in re.findall(r"[a-z0-9]+", topic.lower()) if len(w) > 2]
    kept = [m for m in mem if not (words and any(w in m["text"].lower() for w in words))]
    removed = len(mem) - len(kept)
    _save_json(MEMORY_FILE, kept)
    return f"Forgot {removed} item{'s' if removed != 1 else ''}."


# ---------- reminders (survive restarts; a background thread fires them) ----------

_DAY_ABBR = ["mon", "tue", "wed", "thu", "fri", "sat", "sun"]


def _reminders():
    return _load_json(REMINDERS_FILE, [])


def _save_reminders(r):
    _save_json(REMINDERS_FILE, r)


def _next_id(items):
    return (max([i.get("id", 0) for i in items], default=0)) + 1


def set_reminder(text, at):
    """One-off reminder. `at` is an absolute time 'YYYY-MM-DD HH:MM' (24-hour)."""
    dt = _parse_dt(at)
    if dt is None:
        return f"Couldn't understand the time '{at}'. Use YYYY-MM-DD HH:MM."
    r = _reminders()
    r.append({"id": _next_id(r), "text": text, "type": "once", "at": dt.strftime("%Y-%m-%d %H:%M")})
    _save_reminders(r)
    return f"Reminder set for {dt.strftime('%a %d %b %I:%M %p')}: {text}"


def set_recurring_reminder(text, time, repeat="daily"):
    """Repeating reminder. `time` is 'HH:MM'. `repeat` is 'daily', 'weekdays', 'weekends',
    or a comma list of day abbreviations like 'mon,wed,fri'."""
    if _parse_dt("2000-01-01 " + time) is None:
        return f"Couldn't understand the time '{time}'. Use HH:MM."
    rep = repeat.lower().strip()
    if rep not in ("daily", "weekdays", "weekends"):
        days = [d[:3] for d in re.split(r"[,\s]+", rep) if d[:3] in _DAY_ABBR]
        if not days:
            return "Repeat must be daily, weekdays, weekends, or days like 'mon,wed,fri'."
        rep = ",".join(days)
    hh, mm = time.split(":")
    now = datetime.now()
    due_today = now.replace(hour=int(hh), minute=int(mm), second=0, microsecond=0)
    last = now.strftime("%Y-%m-%d") if now >= due_today else ""   # don't fire retroactively today
    r = _reminders()
    r.append({"id": _next_id(r), "text": text, "type": "recurring",
              "time": time, "repeat": rep, "last": last})
    _save_reminders(r)
    return f"Recurring reminder set for {time} ({rep}): {text}"


def list_reminders():
    """List all active reminders."""
    r = _reminders()
    if not r:
        return "No reminders set."
    out = []
    for i, x in enumerate(r, 1):
        if x["type"] == "once":
            out.append(f"{i}. {x['at']} - {x['text']}")
        else:
            out.append(f"{i}. {x['time']} {x['repeat']} - {x['text']}")
    return "\n".join(out)


def cancel_reminder(which):
    """Cancel a reminder by its number (from list_reminders) or by matching text."""
    r = _reminders()
    if not r:
        return "No reminders to cancel."
    idx = None
    s = str(which).strip()
    if s.isdigit() and 1 <= int(s) <= len(r):
        idx = int(s) - 1
    else:
        for i, x in enumerate(r):
            if s.lower() in x["text"].lower():
                idx = i
                break
    if idx is None:
        return f"Couldn't find a reminder matching '{which}'."
    gone = r.pop(idx)
    _save_reminders(r)
    return f"Cancelled: {gone['text']}"


def _reminder_tick():
    """Called every ~30s by Master. Fires anything due; reschedules recurring ones."""
    r = _reminders()
    if not r:
        return
    now = datetime.now()
    today = now.strftime("%Y-%m-%d")
    changed = False
    keep = []
    for x in r:
        fired = False
        if x["type"] == "once":
            due = _parse_dt(x["at"])
            if due and now >= due:
                if announce:
                    announce(f"Reminder: {x['text']}")
                fired = True
                changed = True
        else:  # recurring
            hh, mm = x["time"].split(":")
            due_today = now.replace(hour=int(hh), minute=int(mm), second=0, microsecond=0)
            wd = now.weekday()  # 0=Mon
            rep = x["repeat"]
            match = (rep == "daily"
                     or (rep == "weekdays" and wd < 5)
                     or (rep == "weekends" and wd >= 5)
                     or (_DAY_ABBR[wd] in rep.split(",")))
            if match and now >= due_today and x.get("last") != today:
                if announce:
                    announce(f"Reminder: {x['text']}")
                x["last"] = today
                changed = True
        if not fired:
            keep.append(x)
    if changed:
        _save_reminders(keep)


# ---------- file reading ----------

_SEARCH_DIRS = [
    os.path.join(os.path.expanduser("~"), d)
    for d in ("Desktop", "Documents", "Downloads", "Pictures")
] + [os.path.expanduser("~")]

_TEXT_EXT = {".txt", ".md", ".csv", ".log", ".json", ".py", ".ini", ".bat", ".xml", ".html"}


def find_file(name):
    """Search Desktop, Documents, Downloads etc. for files whose name contains `name`."""
    name = name.lower().strip()
    hits = []
    for base in _SEARCH_DIRS:
        for path in glob.glob(os.path.join(base, "**", "*"), recursive=True):
            if os.path.isfile(path) and name in os.path.basename(path).lower():
                hits.append(path)
            if len(hits) >= 15:
                break
    if not hits:
        return f"No files found matching '{name}'."
    return "\n".join(hits[:15])


def _resolve(path_or_name):
    if os.path.isfile(path_or_name):
        return path_or_name
    res = find_file(path_or_name)
    first = res.splitlines()[0]
    return first if os.path.isfile(first) else None


def read_file(path_or_name, max_chars=8000):
    """Read a text, PDF, Word, or Excel file and return its text (truncated). Accepts a path or a filename to search for."""
    path = _resolve(path_or_name)
    if not path:
        return f"Couldn't find a file for '{path_or_name}'."
    ext = os.path.splitext(path)[1].lower()
    try:
        if ext == ".pdf":
            text = "\n".join((pg.extract_text() or "") for pg in _PdfReader(path).pages)
        elif ext == ".docx":
            text = "\n".join(p.text for p in _Docx(path).paragraphs)
        elif ext == ".xlsx":
            ws = _load_wb(path, read_only=True).active
            text = "\n".join(
                ", ".join("" if c is None else str(c) for c in row)
                for row in ws.iter_rows(values_only=True))
        elif ext in _TEXT_EXT or ext == "":
            with open(path, "r", encoding="utf-8", errors="replace") as f:
                text = f.read()
        else:
            return f"Can't read {ext} files."
    except Exception as e:
        return f"Couldn't read {os.path.basename(path)}: {e}"
    text = text.strip()
    if len(text) > max_chars:
        text = text[:max_chars] + "\n...[truncated]"
    return f"{os.path.basename(path)}:\n{text}" if text else f"{os.path.basename(path)} is empty."


# ---------- phone (via Windows Phone Link on the paired phone) ----------

def add_contact(name, number):
    """Save a phone number for a name so 'call <name>' works later."""
    contacts = _load_json(CONTACTS_FILE, {})
    contacts[name.strip().lower()] = re.sub(r"[^\d+]", "", number)
    _save_json(CONTACTS_FILE, contacts)
    return f"Saved {name}'s number."


def list_contacts():
    """List saved contacts."""
    contacts = _load_json(CONTACTS_FILE, {})
    if not contacts:
        return "No contacts saved yet."
    return "\n".join(f"{n}: {num}" for n, num in contacts.items())


def call(target):
    """Start a phone call through the linked phone (Windows Phone Link).
    `target` is a saved contact name or a phone number."""
    contacts = _load_json(CONTACTS_FILE, {})
    t = target.strip()
    number = contacts.get(t.lower())
    if number is None and any(c.isdigit() for c in t):
        number = re.sub(r"[^\d+]", "", t)
    if not number:
        return f"I don't have a number for '{target}'. Save it first, e.g. \"save {target}'s number as ...\"."
    try:
        os.startfile(f"tel:{number}")
        who = next((n for n, x in contacts.items() if x == number), number)
        return f"Calling {who}."
    except Exception as e:
        return f"Couldn't start the call: {e}"


# ---------- utilities ----------

_MATH_NS = {k: getattr(math, k) for k in
           ("sqrt", "sin", "cos", "tan", "asin", "acos", "atan", "log", "log10",
            "log2", "exp", "pi", "e", "floor", "ceil", "pow", "fabs", "factorial",
            "radians", "degrees")}
_MATH_NS["abs"] = abs
_MATH_NS["round"] = round


def calculate(expression):
    """Work out a maths expression precisely. Supports + - * / ** %, parentheses and common math functions."""
    expr = expression.strip()
    if len(expr) > 200 or re.search(r"[^0-9a-zA-Z_+\-*/%.,()\s]", expr) or "__" in expr:
        return "That expression has characters I won't evaluate."
    try:
        val = eval(expr, {"__builtins__": {}}, _MATH_NS)   # namespace is locked down
        if isinstance(val, float) and val.is_integer():
            val = int(val)
        return f"{expr} = {val}"
    except Exception as e:
        return f"Couldn't calculate that: {e}"


def clipboard_read():
    """Return whatever text is currently on the clipboard."""
    try:
        out = subprocess.run(["powershell", "-NoProfile", "-Command", "Get-Clipboard -Raw"],
                             capture_output=True, text=True, timeout=10).stdout
        return out.strip() or "The clipboard is empty."
    except Exception as e:
        return f"Couldn't read the clipboard: {e}"


def clipboard_write(text):
    """Copy text onto the clipboard."""
    try:
        subprocess.run(["powershell", "-NoProfile", "-Command",
                        "$in = [Console]::In.ReadToEnd(); Set-Clipboard -Value $in"],
                       input=text, text=True, timeout=10)
        return "Copied to the clipboard."
    except Exception as e:
        return f"Couldn't set the clipboard: {e}"


def system_status():
    """Battery, CPU, memory and disk status of this PC."""
    ps = (
        "$b=(Get-CimInstance Win32_Battery).EstimatedChargeRemaining;"
        "$os=Get-CimInstance Win32_OperatingSystem;"
        "$mu=[math]::Round(($os.TotalVisibleMemorySize-$os.FreePhysicalMemory)/1MB,1);"
        "$mt=[math]::Round($os.TotalVisibleMemorySize/1MB,1);"
        "$cpu=[math]::Round((Get-CimInstance Win32_Processor|Measure-Object -Property LoadPercentage -Average).Average);"
        "$df=[math]::Round((Get-PSDrive C).Free/1GB);"
        "\"battery=$b|cpu=$cpu|mu=$mu|mt=$mt|df=$df\""
    )
    try:
        out = subprocess.run(["powershell", "-NoProfile", "-Command", ps],
                             capture_output=True, text=True, timeout=15).stdout.strip()
        d = dict(p.split("=", 1) for p in out.split("|") if "=" in p)
        batt = f"{d['battery']}% battery" if d.get("battery") else "no battery detected"
        return (f"{batt}, CPU around {d.get('cpu', '?')}%, "
                f"memory {d.get('mu', '?')} of {d.get('mt', '?')} GB in use, "
                f"{d.get('df', '?')} GB free on the C drive.")
    except Exception as e:
        return f"Couldn't read the system status: {e}"


def take_screenshot():
    """Capture the whole screen to an image and open it."""
    try:
        from PIL import ImageGrab
        img = ImageGrab.grab()
        path = _new_path("Screenshot", "png")
        img.save(path)
        os.startfile(path)
        return f"Saved and opened {os.path.basename(path)}."
    except Exception as e:
        return f"Couldn't take a screenshot: {e}"


def convert_currency(amount, from_code, to_code):
    """Convert money between currencies using live rates (no API key)."""
    try:
        a = float(amount)
        fc, tc = from_code.strip().upper(), to_code.strip().upper()
        r = requests.get(f"https://open.er-api.com/v6/latest/{fc}", timeout=12).json()
        rate = (r.get("rates") or {}).get(tc)
        if r.get("result") != "success" or rate is None:
            return f"Couldn't get a rate for {fc} to {tc}."
        return f"{a:g} {fc} = {a * rate:,.2f} {tc}"
    except Exception as e:
        return f"Currency conversion failed: {e}"


# ---------- task list tools ----------

def _load_tasks():
    """Helper: read the task list from disk (empty list if no file yet)."""
    if os.path.exists(TASKS_FILE):
        with open(TASKS_FILE, "r", encoding="utf-8") as f:
            return json.load(f)
    return []


def _save_tasks(tasks):
    """Helper: write the task list to disk."""
    with open(TASKS_FILE, "w", encoding="utf-8") as f:
        json.dump(tasks, f, indent=2, ensure_ascii=False)


def add_task(text):
    """Add a task to the to-do list."""
    tasks = _load_tasks()
    tasks.append({"text": text, "done": False})
    _save_tasks(tasks)
    return f"Added task: {text}"


def list_tasks():
    """Return the whole to-do list as a numbered list."""
    tasks = _load_tasks()
    if not tasks:
        return "The to-do list is empty."
    lines = []
    for i, t in enumerate(tasks, start=1):
        mark = "x" if t["done"] else " "
        lines.append(f"{i}. [{mark}] {t['text']}")
    return "\n".join(lines)


# ---------- the menu the model sees ----------

# Descriptions here are deliberately terse - the full schema is sent on EVERY
# model call, and its token cost eats into API rate limits. Keep only what
# disambiguates the tool; skip "use when the user asks..." filler the model
# already infers from the request itself.
TOOLS = [
    {"type": "function", "function": {
        "name": "get_time", "description": "Current date and time.",
        "parameters": {"type": "object", "properties": {}, "required": []}}},
    {"type": "function", "function": {
        "name": "web_search", "description": "Search the web for current facts, news, or anything you're unsure of.",
        "parameters": {"type": "object", "properties": {
            "query": {"type": "string"}}, "required": ["query"]}}},
    {"type": "function", "function": {
        "name": "get_weather", "description": "Current weather for a place.",
        "parameters": {"type": "object", "properties": {
            "place": {"type": "string"}}, "required": ["place"]}}},
    {"type": "function", "function": {
        "name": "open_website", "description": "Open a URL in the browser.",
        "parameters": {"type": "object", "properties": {
            "url": {"type": "string"}}, "required": ["url"]}}},
    {"type": "function", "function": {
        "name": "open_app", "description": "Launch a Windows program by name.",
        "parameters": {"type": "object", "properties": {
            "name": {"type": "string"}}, "required": ["name"]}}},
    {"type": "function", "function": {
        "name": "media_control", "description": "Media playback / volume keys.",
        "parameters": {"type": "object", "properties": {
            "action": {"type": "string", "enum": ["play_pause", "next", "previous", "stop",
                                                  "volume_up", "volume_down", "mute"]}},
            "required": ["action"]}}},
    {"type": "function", "function": {
        "name": "system_control", "description": "Lock the screen or sleep the PC (only on a clear request).",
        "parameters": {"type": "object", "properties": {
            "action": {"type": "string", "enum": ["lock", "sleep"]}}, "required": ["action"]}}},
    {"type": "function", "function": {
        "name": "set_timer", "description": "Countdown timer; announces itself when done.",
        "parameters": {"type": "object", "properties": {
            "minutes": {"type": "number", "description": "may be fractional, e.g. 0.5"},
            "label": {"type": "string", "description": "optional, what it's for"}},
            "required": ["minutes"]}}},
    {"type": "function", "function": {
        "name": "create_word_document",
        "description": "Write the content yourself, then create a Word doc and open it. Blank line = new paragraph; '# '/'## ' = headings; '- ' = bullets.",
        "parameters": {"type": "object", "properties": {
            "title": {"type": "string"}, "content": {"type": "string", "description": "full body text you wrote"}},
            "required": ["title", "content"]}}},
    {"type": "function", "function": {
        "name": "create_powerpoint",
        "description": "Write the slides yourself, then create a PowerPoint and open it. Each slide: title line, '- ' bullets, optional 'image: <desc>' line; blank line between slides.",
        "parameters": {"type": "object", "properties": {
            "title": {"type": "string"}, "slides": {"type": "string", "description": "slides you wrote, in the described format"}},
            "required": ["title", "slides"]}}},
    {"type": "function", "function": {
        "name": "create_spreadsheet",
        "description": "Write the data yourself, then create an Excel file and open it. One row per line, cells separated by | or comma, header row first.",
        "parameters": {"type": "object", "properties": {
            "title": {"type": "string"}, "data": {"type": "string", "description": "rows you built"}},
            "required": ["title", "data"]}}},
    {"type": "function", "function": {
        "name": "draft_email",
        "description": "Write the body yourself, then open a pre-filled draft in the mail app. Never sends - the user sends it. Keep it short.",
        "parameters": {"type": "object", "properties": {
            "to": {"type": "string", "description": "recipient, or empty if unknown"},
            "subject": {"type": "string"}, "body": {"type": "string"}},
            "required": ["subject", "body"]}}},
    {"type": "function", "function": {
        "name": "list_calendar_events", "description": "List upcoming Outlook calendar events.",
        "parameters": {"type": "object", "properties": {
            "days": {"type": "integer", "description": "days ahead, default 7"}}, "required": []}}},
    {"type": "function", "function": {
        "name": "add_calendar_event",
        "description": "Create a calendar invite to open and accept (not saved automatically). Give an absolute date-time.",
        "parameters": {"type": "object", "properties": {
            "title": {"type": "string"},
            "start": {"type": "string", "description": "'YYYY-MM-DD HH:MM' 24h. For a relative day like 'tomorrow', call get_time first."},
            "minutes": {"type": "integer", "description": "duration, default 60"}},
            "required": ["title", "start"]}}},
    {"type": "function", "function": {
        "name": "remember", "description": "Save a fact about the user for the long term.",
        "parameters": {"type": "object", "properties": {
            "fact": {"type": "string", "description": "as a full sentence"}}, "required": ["fact"]}}},
    {"type": "function", "function": {
        "name": "recall", "description": "Look up a saved fact.",
        "parameters": {"type": "object", "properties": {
            "topic": {"type": "string", "description": "empty = list everything"}}, "required": []}}},
    {"type": "function", "function": {
        "name": "forget", "description": "Delete saved facts matching a topic.",
        "parameters": {"type": "object", "properties": {"topic": {"type": "string"}}, "required": ["topic"]}}},
    {"type": "function", "function": {
        "name": "set_reminder",
        "description": "One-off reminder at an absolute time; speaks up then, survives restarts. For a relative time, call get_time first.",
        "parameters": {"type": "object", "properties": {
            "text": {"type": "string"}, "at": {"type": "string", "description": "'YYYY-MM-DD HH:MM' 24h"}},
            "required": ["text", "at"]}}},
    {"type": "function", "function": {
        "name": "set_recurring_reminder", "description": "Repeating reminder at a time of day.",
        "parameters": {"type": "object", "properties": {
            "text": {"type": "string"}, "time": {"type": "string", "description": "'HH:MM' 24h"},
            "repeat": {"type": "string", "description": "daily / weekdays / weekends / 'mon,wed,fri'"}},
            "required": ["text", "time"]}}},
    {"type": "function", "function": {
        "name": "list_reminders", "description": "List active reminders.",
        "parameters": {"type": "object", "properties": {}, "required": []}}},
    {"type": "function", "function": {
        "name": "cancel_reminder", "description": "Cancel a reminder by number (from list_reminders) or matching text.",
        "parameters": {"type": "object", "properties": {"which": {"type": "string"}}, "required": ["which"]}}},
    {"type": "function", "function": {
        "name": "call", "description": "Call a saved contact or number via the linked phone. Starts a real call.",
        "parameters": {"type": "object", "properties": {"target": {"type": "string"}}, "required": ["target"]}}},
    {"type": "function", "function": {
        "name": "add_contact", "description": "Save a phone number under a name.",
        "parameters": {"type": "object", "properties": {
            "name": {"type": "string"}, "number": {"type": "string", "description": "ideally with country code"}},
            "required": ["name", "number"]}}},
    {"type": "function", "function": {
        "name": "list_contacts", "description": "List saved phone contacts.",
        "parameters": {"type": "object", "properties": {}, "required": []}}},
    {"type": "function", "function": {
        "name": "find_file", "description": "Find a file by name in Desktop/Documents/Downloads/home.",
        "parameters": {"type": "object", "properties": {"name": {"type": "string"}}, "required": ["name"]}}},
    {"type": "function", "function": {
        "name": "read_file",
        "description": "Read a text/PDF/Word/Excel file's contents (to summarise or read aloud). Path or filename to search for.",
        "parameters": {"type": "object", "properties": {"path_or_name": {"type": "string"}}, "required": ["path_or_name"]}}},
    {"type": "function", "function": {
        "name": "add_task", "description": "Add an item to the to-do list.",
        "parameters": {"type": "object", "properties": {"text": {"type": "string"}}, "required": ["text"]}}},
    {"type": "function", "function": {
        "name": "list_tasks", "description": "Show the to-do list.",
        "parameters": {"type": "object", "properties": {}, "required": []}}},
    {"type": "function", "function": {
        "name": "calculate", "description": "Evaluate a maths expression exactly (use instead of doing arithmetic yourself).",
        "parameters": {"type": "object", "properties": {
            "expression": {"type": "string", "description": "e.g. '12.5*(3+4)/2' or 'sqrt(144)'"}}, "required": ["expression"]}}},
    {"type": "function", "function": {
        "name": "clipboard_read", "description": "Read the current clipboard text.",
        "parameters": {"type": "object", "properties": {}, "required": []}}},
    {"type": "function", "function": {
        "name": "clipboard_write", "description": "Put text on the clipboard.",
        "parameters": {"type": "object", "properties": {"text": {"type": "string"}}, "required": ["text"]}}},
    {"type": "function", "function": {
        "name": "system_status", "description": "Battery, CPU, memory and disk status of this PC.",
        "parameters": {"type": "object", "properties": {}, "required": []}}},
    {"type": "function", "function": {
        "name": "take_screenshot", "description": "Screenshot the screen, save it, and open it.",
        "parameters": {"type": "object", "properties": {}, "required": []}}},
    {"type": "function", "function": {
        "name": "convert_currency", "description": "Convert money between currencies at live rates.",
        "parameters": {"type": "object", "properties": {
            "amount": {"type": "number"}, "from_code": {"type": "string", "description": "e.g. USD"},
            "to_code": {"type": "string", "description": "e.g. INR"}},
            "required": ["amount", "from_code", "to_code"]}}},
]

TOOL_FUNCTIONS = {
    "get_time": get_time,
    "web_search": web_search,
    "get_weather": get_weather,
    "open_website": open_website,
    "open_app": open_app,
    "media_control": media_control,
    "system_control": system_control,
    "set_timer": set_timer,
    "create_word_document": create_word_document,
    "create_powerpoint": create_powerpoint,
    "create_spreadsheet": create_spreadsheet,
    "draft_email": draft_email,
    "list_calendar_events": list_calendar_events,
    "add_calendar_event": add_calendar_event,
    "remember": remember,
    "recall": recall,
    "forget": forget,
    "set_reminder": set_reminder,
    "set_recurring_reminder": set_recurring_reminder,
    "list_reminders": list_reminders,
    "cancel_reminder": cancel_reminder,
    "call": call,
    "add_contact": add_contact,
    "list_contacts": list_contacts,
    "find_file": find_file,
    "read_file": read_file,
    "add_task": add_task,
    "list_tasks": list_tasks,
    "calculate": calculate,
    "clipboard_read": clipboard_read,
    "clipboard_write": clipboard_write,
    "system_status": system_status,
    "take_screenshot": take_screenshot,
    "convert_currency": convert_currency,
}
